from flask import (
    Blueprint, render_template, request, redirect, url_for,
    flash, jsonify, Response, send_file
)
from datetime import datetime
import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from extensions import db
from models import Venta, VentaDetalle, Producto, Inventario, Cliente


ventas_bp = Blueprint('ventas', __name__, url_prefix='/ventas')


# =====================================================================
# HELPER: construir reporte (reutilizable por vista y exports)
# =====================================================================
def _construir_reporte(filtro_mes):
    """
    Devuelve (tickets_flat, totales_dict) para reutilizar en vistas y exports.

    tickets_flat: lista de dicts, uno por venta/ticket
    totales_dict: {
        'periodos': { 'YYYY-MM': [tickets...] },
        'total_ingresos': float,
        'total_ganancia': float,
        'total_tickets': int,
    }
    """
    query = Venta.query.order_by(Venta.fecha.desc(), Venta.id.desc())

    if filtro_mes:
        try:
            partes = filtro_mes.split('-')
            if len(partes) == 2:
                anio = int(partes[0])
                mes = int(partes[1])
                if 1 <= mes <= 12:
                    query = query.filter(
                        db.extract('year', Venta.fecha) == anio,
                        db.extract('month', Venta.fecha) == mes
                    )
        except (ValueError, TypeError):
            pass

    ventas = query.all()

    # ---- Detalles en batch (evita N+1) ----
    venta_ids = [v.id for v in ventas]
    detalles_por_venta = {}
    if venta_ids:
        todos_detalles = (
            VentaDetalle.query
            .filter(VentaDetalle.id_venta.in_(venta_ids))
            .all()
        )
        for d in todos_detalles:
            detalles_por_venta.setdefault(d.id_venta, []).append(d)

    # ---- Productos en batch (evita N+1) ----
    producto_ids = {d.id_producto for d in sum(detalles_por_venta.values(), [])}
    productos_map = {}
    if producto_ids:
        productos = Producto.query.filter(Producto.id.in_(producto_ids)).all()
        productos_map = {p.id: p for p in productos}

    # ---- Construir estructura de tickets ----
    tickets = []
    for v in ventas:
        detalles_venta = detalles_por_venta.get(v.id, [])
        detalles = []
        total_costo = 0.0

        for d in detalles_venta:
            producto = productos_map.get(d.id_producto)
            cant = float(d.cantidad or 0)
            precio_u = float(d.precio_unitario or 0)
            costo_u = float(d.costo_unitario or 0)
            subtotal = cant * precio_u
            costo_linea = cant * costo_u
            ganancia_linea = subtotal - costo_linea
            total_costo += costo_linea

            detalles.append({
                'sku': producto.sku if producto else 'N/A',
                'nombre': producto.nombre if producto else 'Producto eliminado',
                'cantidad': d.cantidad,
                'precio_unitario': precio_u,
                'costo_unitario': costo_u,
                'subtotal': subtotal,
                'ganancia_linea': ganancia_linea,
            })

        total = float(v.total or 0)

        tickets.append({
            'folio': v.folio,
            'fecha': v.fecha,
            'anio': v.fecha.year if v.fecha else None,
            'mes': v.fecha.month if v.fecha else None,
            'cliente': v.cliente.nombre if v.cliente else 'Público general',
            'canal': v.canal or '—',
            'estatus': v.estatus or '—',
            'subtotal': float(v.subtotal or 0),
            'descuento': float(v.descuento or 0),
            'total': total,
            'total_costo': total_costo,
            'ganancia': total - total_costo,
            'detalles': detalles,
        })

    # ---- Agrupar por periodo ----
    periodos = {}
    for t in tickets:
        if t['anio'] and t['mes']:
            key = f"{t['anio']}-{t['mes']:02d}"
            periodos.setdefault(key, []).append(t)

    totales = {
        'periodos': periodos,
        'total_ingresos': sum(t['total'] for t in tickets),
        'total_ganancia': sum(t['ganancia'] for t in tickets),
        'total_tickets': len(tickets),
    }

    return tickets, totales


# =====================================================================
# VENTAS: Listado
# =====================================================================
@ventas_bp.route('/')
def index():
    filtro_dia = request.args.get('dia', '')
    filtro_mes = request.args.get('mes', '')
    filtro_canal = request.args.get('canal', '')

    query = Venta.query

    if filtro_dia:
        try:
            fecha_inicio = datetime.strptime(filtro_dia, '%Y-%m-%d')
            fecha_fin = datetime.strptime(filtro_dia + ' 23:59:59', '%Y-%m-%d %H:%M:%S')
            query = query.filter(Venta.fecha >= fecha_inicio, Venta.fecha <= fecha_fin)
        except ValueError:
            pass

    if filtro_mes and not filtro_dia:
        try:
            partes = filtro_mes.split('-')
            if len(partes) == 2:
                anio = int(partes[0])
                mes = int(partes[1])
                if 1 <= mes <= 12:
                    query = query.filter(
                        db.extract('year', Venta.fecha) == anio,
                        db.extract('month', Venta.fecha) == mes
                    )
        except (ValueError, TypeError):
            pass

    if filtro_canal:
        query = query.filter(Venta.canal.ilike(f"%{filtro_canal}%"))

    ventas = query.order_by(Venta.id.desc()).all()

    return render_template(
        'ventas/index.html',
        ventas=ventas,
        filtro_dia=filtro_dia,
        filtro_mes=filtro_mes,
        filtro_canal=filtro_canal
    )


# =====================================================================
# VENTAS: Nueva
# =====================================================================
@ventas_bp.route('/nueva', methods=['GET', 'POST'])
def nueva():
    if request.method == 'POST':
        canal = request.form.get('canal')
        id_cliente = request.form.get('id_cliente')
        fecha_str = request.form.get('fecha')

        fecha_venta = datetime.strptime(fecha_str, '%Y-%m-%d').date() if fecha_str else datetime.utcnow().date()

        productos_ids = request.form.getlist('producto_id[]')
        cantidades = request.form.getlist('cantidad[]')
        precios = request.form.getlist('precio[]')

        if not productos_ids:
            flash('Debe agregar al menos un producto a la venta.', 'danger')
            return redirect(url_for('ventas.nueva'))

        total_venta = 0
        detalles_temp = []

        for i in range(len(productos_ids)):
            p_id = int(productos_ids[i])
            cant = int(cantidades[i])
            precio_u = float(precios[i])
            subtotal = cant * precio_u
            total_venta += subtotal

            producto_db = Producto.query.get(p_id)
            costo_u = producto_db.costo if hasattr(producto_db, 'costo') and producto_db.costo else 0.0

            detalles_temp.append({
                'id_producto': p_id,
                'cantidad': cant,
                'precio_unitario': precio_u,
                'costo_unitario': costo_u
            })

        ultimo_id = db.session.query(db.func.max(Venta.id)).scalar() or 0
        folio_automatico = f"F-{ultimo_id + 1:05d}"

        nueva_venta = Venta(
            folio=folio_automatico,
            canal=canal,
            id_cliente=int(id_cliente) if id_cliente else None,
            fecha=fecha_venta,
            total=total_venta,
            subtotal=total_venta,
            descuento=0.0,
            estatus='COMPLETADA'
        )
        db.session.add(nueva_venta)
        db.session.commit()

        for det in detalles_temp:
            detalle_venta = VentaDetalle(
                id_venta=nueva_venta.id,
                id_producto=det['id_producto'],
                cantidad=det['cantidad'],
                precio_unitario=det['precio_unitario'],
                costo_unitario=det['costo_unitario'],
                descuento=0.0
            )
            db.session.add(detalle_venta)

            inv = Inventario.query.filter_by(id_producto=det['id_producto'], id_ubicacion=1).first()
            if inv:
                inv.existencia -= det['cantidad']
                if inv.existencia < 0:
                    inv.existencia = 0

        db.session.commit()
        flash('¡Venta registrada con éxito!', 'success')
        return redirect(url_for('ventas.recibo', id_venta=nueva_venta.id))

    clientes = Cliente.query.all()
    productos_stock = db.session.query(Producto, Inventario.existencia)\
        .join(Inventario, Producto.id == Inventario.id_producto)\
        .filter(Inventario.id_ubicacion == 1, Producto.activo == 1).all()

    fecha_hoy = datetime.now().strftime('%Y-%m-%d')
    return render_template('ventas/nueva.html', clientes=clientes, productos=productos_stock, fecha_hoy=fecha_hoy)


# =====================================================================
# VENTAS: Editar
# =====================================================================
@ventas_bp.route('/editar/<int:id>', methods=['GET', 'POST'])
def editar(id):
    venta = Venta.query.get_or_404(id)

    if request.method == 'POST':
        canal = request.form.get('canal')
        id_cliente = request.form.get('id_cliente')
        fecha_str = request.form.get('fecha')

        if fecha_str:
            try:
                venta.fecha = datetime.strptime(fecha_str, '%Y-%m-%d')
            except ValueError:
                pass

        if canal:
            venta.canal = canal

        venta.id_cliente = int(id_cliente) if id_cliente else None

        # Reintegrar inventario anterior
        detalles_actuales = VentaDetalle.query.filter_by(id_venta=venta.id).all()
        for det_viejo in detalles_actuales:
            inv = Inventario.query.filter_by(id_producto=det_viejo.id_producto, id_ubicacion=1).first()
            if inv:
                inv.existencia += det_viejo.cantidad
            db.session.delete(det_viejo)

        productos_ids = request.form.getlist('producto_id[]')
        cantidades = request.form.getlist('cantidad[]')
        precios = request.form.getlist('precio[]')

        if not productos_ids:
            flash('La venta debe tener al menos un producto.', 'danger')
            return redirect(url_for('ventas.editar', id=venta.id))

        total_venta = 0
        for i in range(len(productos_ids)):
            p_id = int(productos_ids[i])
            cant = int(cantidades[i])
            precio_u = float(precios[i])
            subtotal = cant * precio_u
            total_venta += subtotal

            producto_db = Producto.query.get(p_id)
            costo_u = producto_db.costo if hasattr(producto_db, 'costo') and producto_db.costo else 0.0

            nuevo_detalle = VentaDetalle(
                id_venta=venta.id,
                id_producto=p_id,
                cantidad=cant,
                precio_unitario=precio_u,
                costo_unitario=costo_u,
                descuento=0.0
            )
            db.session.add(nuevo_detalle)

            inv = Inventario.query.filter_by(id_producto=p_id, id_ubicacion=1).first()
            if inv:
                inv.existencia -= cant
                if inv.existencia < 0:
                    inv.existencia = 0

        venta.total = total_venta
        venta.subtotal = total_venta

        db.session.commit()
        flash('¡Venta actualizada con éxito!', 'success')
        return redirect(url_for('ventas.recibo', id_venta=venta.id))

    clientes = Cliente.query.all()
    productos_stock = db.session.query(Producto, Inventario.existencia)\
        .join(Inventario, Producto.id == Inventario.id_producto)\
        .filter(Inventario.id_ubicacion == 1, Producto.activo == 1).all()

    detalles = VentaDetalle.query.options(db.joinedload(VentaDetalle.producto)).filter_by(id_venta=venta.id).all()

    return render_template('ventas/editar.html', venta=venta, clientes=clientes, productos=productos_stock, detalles=detalles)


# =====================================================================
# CLIENTES: Alta rápida vía AJAX
# =====================================================================
@ventas_bp.route('/cliente-rapido', methods=['POST'])
def cliente_rapido():
    nombre = request.form.get('nombre')
    telefono = request.form.get('telefono')
    email = request.form.get('email')

    if not nombre:
        return jsonify({'success': False, 'message': 'El nombre es obligatorio'})

    try:
        nuevo_cliente = Cliente(nombre=nombre, telefono=telefono, email=email)
        db.session.add(nuevo_cliente)
        db.session.commit()
        return jsonify({
            'success': True,
            'cliente': {
                'id': nuevo_cliente.id,
                'nombre': nuevo_cliente.nombre,
                'telefono': nuevo_cliente.telefono
            }
        })
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)})


# =====================================================================
# VENTAS: Recibo
# =====================================================================
@ventas_bp.route('/recibo/<int:id_venta>')
def recibo(id_venta):
    venta = Venta.query.get_or_404(id_venta)
    detalles = VentaDetalle.query.options(db.joinedload(VentaDetalle.producto)).filter_by(id_venta=id_venta).all()
    return render_template('ventas/recibo.html', venta=venta, detalles=detalles)


# =====================================================================
# REPORTE MENSUAL (vista HTML)
# =====================================================================
@ventas_bp.route('/reporte-mensual')
def reporte_mensual():
    filtro_mes = request.args.get('mes', '')
    _, totales = _construir_reporte(filtro_mes)

    return render_template(
        'ventas/reporte_mensual.html',
        periodos=totales['periodos'],
        filtro_mes=filtro_mes,
        total_ingresos=totales['total_ingresos'],
        total_ganancia=totales['total_ganancia'],
        total_tickets=totales['total_tickets'],
    )


# =====================================================================
# EXPORTACIÓN: CSV
# =====================================================================
@ventas_bp.route('/reporte-mensual/exportar/csv')
def reporte_mensual_csv():
    """Exporta el reporte mensual a CSV (una fila por línea de producto)."""
    filtro_mes = request.args.get('mes', '')
    _, totales = _construir_reporte(filtro_mes)

    output = io.StringIO()
    writer = csv.writer(output)

    writer.writerow([
        'Periodo', 'Folio', 'Fecha', 'Cliente', 'Canal', 'Estatus',
        'SKU', 'Producto', 'Cantidad', 'Precio Unitario',
        'Costo Unitario', 'Subtotal', 'Ganancia Línea',
        'Total Ticket', 'Costo Ticket', 'Ganancia Ticket'
    ])

    for periodo, tkts in sorted(totales['periodos'].items(), reverse=True):
        for t in tkts:
            fecha_str = t['fecha'].strftime('%Y-%m-%d') if t['fecha'] else ''
            for d in t['detalles']:
                writer.writerow([
                    periodo,
                    t['folio'],
                    fecha_str,
                    t['cliente'],
                    t['canal'],
                    t['estatus'],
                    d['sku'],
                    d['nombre'],
                    d['cantidad'],
                    f"{d['precio_unitario']:.2f}",
                    f"{d['costo_unitario']:.2f}",
                    f"{d['subtotal']:.2f}",
                    f"{d['ganancia_linea']:.2f}",
                    f"{t['total']:.2f}",
                    f"{t['total_costo']:.2f}",
                    f"{t['ganancia']:.2f}",
                ])

    nombre = f"reporte_mensual_{filtro_mes or 'todos'}.csv"
    csv_data = '\ufeff' + output.getvalue()  # BOM para Excel + UTF-8

    return Response(
        csv_data,
        mimetype='text/csv; charset=utf-8',
        headers={'Content-Disposition': f'attachment; filename={nombre}'}
    )


# =====================================================================
# EXPORTACIÓN: EXCEL
# =====================================================================
@ventas_bp.route('/reporte-mensual/exportar/excel')
def reporte_mensual_excel():
    """Exporta el reporte mensual a Excel con formato profesional."""
    filtro_mes = request.args.get('mes', '')
    _, totales = _construir_reporte(filtro_mes)

    wb = Workbook()

    # ---- Estilos compartidos ----
    titulo_font = Font(bold=True, size=14, color="FFFFFF")
    titulo_fill = PatternFill("solid", fgColor="4F46E5")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0F172A")
    total_fill = PatternFill("solid", fgColor="EEF2FF")
    border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB'),
    )
    center = Alignment(horizontal='center', vertical='center')
    right = Alignment(horizontal='right', vertical='center')
    money_format = '"$"#,##0.00'

    # =================================================================
    # HOJA 1: RESUMEN GENERAL
    # =================================================================
    ws_resumen = wb.active
    ws_resumen.title = "Resumen"

    ws_resumen.merge_cells('A1:D1')
    ws_resumen['A1'] = f"Reporte Mensual de Ventas — {filtro_mes or 'Todos los periodos'}"
    ws_resumen['A1'].font = titulo_font
    ws_resumen['A1'].fill = titulo_fill
    ws_resumen['A1'].alignment = center
    ws_resumen.row_dimensions[1].height = 28

    kpis = [
        ("Tickets totales", totales['total_tickets'], False),
        ("Ingresos totales", totales['total_ingresos'], True),
        ("Ganancia estimada", totales['total_ganancia'], True),
    ]
    for idx, (label, value, is_money) in enumerate(kpis, start=3):
        ws_resumen[f'A{idx}'] = label
        ws_resumen[f'B{idx}'] = value
        ws_resumen[f'A{idx}'].font = Font(bold=True)
        ws_resumen[f'B{idx}'].font = Font(bold=True, color="10B981")
        ws_resumen[f'A{idx}'].fill = total_fill
        ws_resumen[f'B{idx}'].fill = total_fill
        if is_money:
            ws_resumen[f'B{idx}'].number_format = money_format

    ws_resumen.column_dimensions['A'].width = 22
    ws_resumen.column_dimensions['B'].width = 18

    # =================================================================
    # HOJA 2: DETALLE POR TICKET
    # =================================================================
    ws = wb.create_sheet("Detalle")

    headers = [
        'Periodo', 'Folio', 'Fecha', 'Cliente', 'Canal', 'Estatus',
        'SKU', 'Producto', 'Cantidad', 'P. Unit.', 'Costo Unit.',
        'Subtotal', 'Ganancia', 'Total Ticket'
    ]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[1].height = 22

    fila = 2
    for periodo, tkts in sorted(totales['periodos'].items(), reverse=True):
        for t in tkts:
            fecha_str = t['fecha'].strftime('%Y-%m-%d') if t['fecha'] else ''
            for d in t['detalles']:
                datos = [
                    periodo, t['folio'], fecha_str, t['cliente'], t['canal'], t['estatus'],
                    d['sku'], d['nombre'], d['cantidad'],
                    d['precio_unitario'], d['costo_unitario'],
                    d['subtotal'], d['ganancia_linea'], t['total']
                ]
                for col, val in enumerate(datos, 1):
                    cell = ws.cell(row=fila, column=col, value=val)
                    cell.border = border
                    if col in (10, 11, 12, 13, 14):  # monetarias
                        cell.number_format = money_format
                        cell.alignment = right
                    elif col == 9:  # cantidad
                        cell.alignment = center
                fila += 1

    ws.freeze_panes = "A2"

    # Auto-ancho
    for col_idx, header in enumerate(headers, 1):
        max_len = len(str(header))
        for row_idx in range(2, fila):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 40)

    # =================================================================
    # HOJA 3: POR PERIODO
    # =================================================================
    ws_per = wb.create_sheet("Por Periodo")
    headers_per = ['Periodo', 'Tickets', 'Ingresos', 'Ganancia']

    for col, h in enumerate(headers_per, 1):
        cell = ws_per.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    fila = 2
    for periodo in sorted(totales['periodos'].keys(), reverse=True):
        tkts = totales['periodos'][periodo]
        ingresos = sum(t['total'] for t in tkts)
        ganancia = sum(t['ganancia'] for t in tkts)

        datos = [periodo, len(tkts), ingresos, ganancia]
        for col, val in enumerate(datos, 1):
            cell = ws_per.cell(row=fila, column=col, value=val)
            cell.border = border
            if col >= 3:
                cell.number_format = money_format
                cell.alignment = right
            elif col == 2:
                cell.alignment = center
        fila += 1

    for col_idx, header in enumerate(headers_per, 1):
        ws_per.column_dimensions[get_column_letter(col_idx)].width = max(len(header) + 4, 15)

    # =================================================================
    # GUARDAR EN MEMORIA Y DEVOLVER
    # =================================================================
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    nombre = f"reporte_mensual_{filtro_mes or 'todos'}.xlsx"

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=nombre
    )